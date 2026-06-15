import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Создаем Excel файл с тестовыми данными HR-статусов
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "HR-статусы обучения"

# Заголовки - import_hr_status ожидает значения по индексам 1-5
# values[1] = ФИО, values[2] = Должность, values[3] = Категория, values[4] = Организация, values[5] = Статус
headers = ['№', 'ФИО', 'Должность', 'Категория', 'Организация', 'Статус обучения']
ws.append(headers)

# Стиль заголовков
header_fill = PatternFill(start_color="AA141E", end_color="AA141E", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Тестовые данные
organizations = [
    'ДГИ / Управление реновации',
    'ДГИ / Управление эксплуатации',
    'ДГИ / Управление кадрами',
    'ДГИ / Юридический отдел',
    'ДГИ / Финансовый отдел',
    'ДГИ / IT-отдел',
    'ДГИ / Отдел документооборота',
    'ДГИ / Отдел безопасности',
]

positions = [
    'Начальник управления',
    'Заместитель начальника',
    'Главный специалист',
    'Ведущий специалист',
    'Специалист',
    'Менеджер',
    'Юрисконсульт',
    'Бухгалтер',
    'Системный администратор',
]

categories = ['Руководитель', 'Специалист', 'Специалист', 'Специалист', 'Специалист', 'Руководитель', 'Специалист', 'Специалист', 'Специалист']

statuses = ['Пройдено', 'Пройдено', 'Пройдено', 'Не пройдено', 'Пройдено', 'Пройдено', 'Не пройдено', 'Пройдено', 'Пройдено']

import random

# Генерируем 50 записей
for i in range(50):
    org_idx = i % len(organizations)
    pos_idx = i % len(positions)
    status_idx = i % len(statuses)
    
    # Добавляем вариативность в статусы
    if i % 5 == 0:
        status = 'Не пройдено'
    else:
        status = 'Пройдено'
    
    fio = f"Сотрудник {i + 1} {random.choice(['Иванов', 'Петров', 'Сидоров', 'Козлов', 'Новиков'])}"
    
    ws.append([
        i + 1,  # №
        fio,
        positions[pos_idx],
        categories[pos_idx],
        organizations[org_idx],
        status
    ])

# Настройка ширины колонок
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 15
ws.column_dimensions['E'].width = 35
ws.column_dimensions['F'].width = 15

# Сохраняем файл
output_path = './hr_status_test.xlsx'
wb.save(output_path)

print(f"Тестовый Excel файл создан: {output_path}")
print("Файл содержит 50 записей с данными HR-статусов обучения")
print("Загрузите этот файл на странице HR-статусы для проверки")
