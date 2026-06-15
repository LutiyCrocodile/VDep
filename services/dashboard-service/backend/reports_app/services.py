import os
import io
import tempfile
from datetime import datetime
from django.conf import settings
import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


class PPTXGenerator:
    """Генератор PowerPoint-презентаций в корпоративном стиле ДГИ"""

    # Корпоративные цвета ДГИ
    PRIMARY_COLOR = RGBColor(0xAA, 0x14, 0x1E)  # Красный ДГИ
    PRIMARY_DARK = RGBColor(0x8B, 0x10, 0x19)  # Тёмно-красный
    ACCENT_COLOR = RGBColor(0xD4, 0x40, 0x4A)  # Светло-красный
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
    DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
    FONT_NAME = 'Century Gothic'

    def __init__(self, report):
        self.report = report
        self.config = report.data if report.data else {}
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def generate(self):
        """Генерирует PPTX-файл и возвращает путь к нему"""
        slides_config = self.config.get('slides', [])

        # Если нет слайдов в конфиге, создаем базовый слайд с информацией об отчёте
        if not slides_config:
            self._create_basic_slide()
        else:
            for slide_config in slides_config:
                layout = slide_config.get('layout', 'content')

                if layout == 'title':
                    self._create_title_slide(slide_config)
                elif layout == 'content':
                    self._create_content_slide(slide_config)
                elif layout == 'two_columns':
                    self._create_two_columns_slide(slide_config)
                elif layout == 'blank':
                    self._create_blank_slide(slide_config)
                else:
                    self._create_content_slide(slide_config)

        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pptx')
        self.prs.save(temp_file.name)
        return temp_file.name

    def _create_title_slide(self, config):
        """Титульный слайд"""
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Тёмно-красный фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.PRIMARY_DARK

        # Декоративная линия
        left = Inches(1)
        top = Inches(2.5)
        width = Inches(3)
        height = Inches(0.05)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.ACCENT_COLOR
        shape.line.fill.background()

        # Заголовок
        title_text = self._get_element_text(config, 'title', 'Отчёт')
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.7), Inches(11), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = self.WHITE
        p.font.name = self.FONT_NAME

        # Подзаголовок
        subtitle_text = self._get_element_text(config, 'subtitle', 'Департамент городского имущества города Москвы')
        txBox = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        p.font.name = self.FONT_NAME

        # Дата
        date_text = self._get_element_text(config, 'date', datetime.now().strftime('%d.%m.%Y'))
        txBox = slide.shapes.add_textbox(Inches(1), Inches(5.0), Inches(11), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = date_text
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.font.name = self.FONT_NAME

    def _create_content_slide(self, config):
        """Слайд с контентом"""
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Белый фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.WHITE

        # Верхняя полоса
        left = Inches(0)
        top = Inches(0)
        width = self.prs.slide_width
        height = Inches(0.08)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.PRIMARY_COLOR
        shape.line.fill.background()

        # Заголовок слайда
        heading = self._get_element_text(config, 'heading', '')
        if heading:
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = heading
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = self.PRIMARY_COLOR
            p.font.name = self.FONT_NAME

        # Элементы слайда
        y_position = Inches(1.3)
        elements = config.get('elements', [])

        for element in elements:
            elem_type = element.get('type')

            if elem_type == 'text' or elem_type == 'heading':
                y_position = self._add_text_element(slide, element, y_position)
            elif elem_type == 'chart':
                y_position = self._add_chart_element(slide, element, y_position)
            elif elem_type == 'table':
                y_position = self._add_table_element(slide, element, y_position)

        # Подвал
        self._add_footer(slide, config.get('id', 1))

    def _create_two_columns_slide(self, config):
        """Слайд с двумя колонками"""
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Белый фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.WHITE

        # Верхняя полоса
        left = Inches(0)
        top = Inches(0)
        width = self.prs.slide_width
        height = Inches(0.08)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.PRIMARY_COLOR
        shape.line.fill.background()

        # Заголовок
        heading = self._get_element_text(config, 'heading', '')
        if heading:
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.7))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = heading
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = self.PRIMARY_COLOR
            p.font.name = self.FONT_NAME

        elements = config.get('elements', [])
        mid = len(elements) // 2 if len(elements) > 1 else 1
        left_elements = elements[:mid]
        right_elements = elements[mid:]

        # Левая колонка
        y_left = Inches(1.3)
        for element in left_elements:
            y_left = self._add_chart_or_text(slide, element, y_left, left_offset=Inches(0.5))

        # Правая колонка
        y_right = Inches(1.3)
        for element in right_elements:
            y_right = self._add_chart_or_text(slide, element, y_right, left_offset=Inches(6.8))

        self._add_footer(slide, config.get('id', 1))

    def _create_blank_slide(self, config):
        """Пустой слайд"""
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        elements = config.get('elements', [])
        y_position = Inches(0.5)

        for element in elements:
            if element.get('type') == 'text':
                y_position = self._add_text_element(slide, element, y_position)

    def _add_text_element(self, slide, element, y_position, left_offset=None):
        """Текстовый блок"""
        if left_offset is None:
            left_offset = Inches(0.8)

        content = element.get('content', '')
        max_width = Inches(5.5) if left_offset > Inches(1) else Inches(11.5)

        txBox = slide.shapes.add_textbox(left_offset, y_position, max_width, Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(14)
        p.font.color.rgb = self.DARK_TEXT
        p.font.name = self.FONT_NAME

        return y_position + Inches(1.5)

    def _add_chart_element(self, slide, element, y_position, left_offset=None):
        """График"""
        if left_offset is None:
            left_offset = Inches(0.8)

        chart_width = Inches(5.5) if left_offset > Inches(1) else Inches(11.5)
        chart_height = Inches(4)

        chart_type = element.get('chart_type', 'bar')
        title = element.get('title', 'График')
        
        # Получаем данные из элемента
        chart_data = element.get('data', {})
        categories = chart_data.get('categories', [])
        values = chart_data.get('values', [])

        fig, ax = plt.subplots(figsize=(8, 4.5))

        if chart_type == 'bar':
            if categories and values:
                ax.bar(categories, values, color=['#AA141E', '#8B1019', '#D4404A', '#666666', '#333333'][:len(categories)])
            else:
                # Fallback если нет данных
                categories = ['ЦАО', 'САО', 'ЮАО', 'ВАО', 'ЗАО']
                values = [1250, 980, 2100, 1500, 1800]
                ax.bar(categories, values, color=['#AA141E', '#8B1019', '#D4404A', '#666666', '#333333'])
        elif chart_type == 'line':
            if categories and values:
                ax.plot(categories, values, color='#AA141E', linewidth=3, marker='o')
                ax.fill_between(range(len(values)), values, alpha=0.3, color='#D4404A')
            else:
                # Fallback если нет данных
                months = ['Янв', 'Фев', 'Мар', 'Апр', 'Май', 'Июн']
                values = [100, 250, 400, 800, 1200, 1500]
                ax.plot(months, values, color='#AA141E', linewidth=3, marker='o')
                ax.fill_between(range(len(months)), values, alpha=0.3, color='#D4404A')
        elif chart_type == 'pie':
            if categories and values:
                ax.pie(values, labels=categories, colors=['#AA141E', '#8B1019', '#D4404A', '#666666'][:len(categories)], autopct='%1.1f%%')
            else:
                # Fallback если нет данных
                labels = ['Реновация', 'Кадры', 'Обращения', 'Прочее']
                sizes = [45, 25, 20, 10]
                colors = ['#AA141E', '#8B1019', '#D4404A', '#666666']
                ax.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%')

        ax.set_title(title, fontsize=14, fontweight='bold', color='#AA141E')

        temp_img = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
        plt.tight_layout()
        plt.savefig(temp_img.name, dpi=150, bbox_inches='tight', facecolor='white')
        plt.close('all')

        slide.shapes.add_picture(temp_img.name, left_offset, y_position, chart_width, chart_height)

        try:
            os.unlink(temp_img.name)
        except:
            pass

        return y_position + chart_height + Inches(0.3)

    def _add_table_element(self, slide, element, y_position, left_offset=None):
        """Таблица"""
        if left_offset is None:
            left_offset = Inches(0.8)

        rows = 5
        cols = 3
        table_width = Inches(5.5) if left_offset > Inches(1) else Inches(11.5)
        table_height = Inches(2.5)

        table_shape = slide.shapes.add_table(rows, cols, left_offset, y_position, table_width, table_height)
        table = table_shape.table

        data = [
            ['Район', 'Показатель', 'Значение'],
            ['ЦАО', 'Переселено семей', '1 250'],
            ['САО', 'Переселено семей', '980'],
            ['ЮАО', 'Переселено семей', '2 100'],
            ['ВАО', 'Переселено семей', '1 500'],
        ]

        for i, row_data in enumerate(data):
            for j, cell_text in enumerate(row_data):
                cell = table.cell(i, j)
                cell.text = cell_text
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(11)
                    paragraph.font.name = self.FONT_NAME
                    if i == 0:
                        paragraph.font.bold = True
                        paragraph.font.color.rgb = self.WHITE
                    else:
                        paragraph.font.color.rgb = self.DARK_TEXT
                if i == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.PRIMARY_COLOR

        return y_position + table_height + Inches(0.3)

    def _add_chart_or_text(self, slide, element, y_position, left_offset):
        """Добавляет элемент в колонку"""
        if element.get('type') == 'chart':
            return self._add_chart_element(slide, element, y_position, left_offset)
        elif element.get('type') == 'table':
            return self._add_table_element(slide, element, y_position, left_offset)
        else:
            return self._add_text_element(slide, element, y_position, left_offset)

    def _add_footer(self, slide, slide_number):
        """Подвал слайда"""
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12), Inches(0.4))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"ДГИ | Отчёт | {datetime.now().strftime('%d.%m.%Y')} | Слайд {slide_number}"
        p.font.size = Pt(8)
        p.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.alignment = PP_ALIGN.RIGHT
        p.font.name = self.FONT_NAME

    def _get_element_text(self, config, elem_type, default=''):
        """Ищет текст элемента в конфигурации"""
        elements = config.get('elements', [])
        for elem in elements:
            if elem.get('type') == elem_type:
                return elem.get('content', default)
        return default

    def _create_basic_slide(self):
        """Создает базовый слайд с информацией об отчёте"""
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Тёмно-красный фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.PRIMARY_DARK

        # Декоративная линия
        left = Inches(1)
        top = Inches(2.5)
        width = Inches(3)
        height = Inches(0.05)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.ACCENT_COLOR
        shape.line.fill.background()

        # Заголовок
        title_text = self.report.title or 'Отчёт'
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.7), Inches(11), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = self.WHITE
        p.font.name = self.FONT_NAME

        # Дата
        date_text = self.report.created_at.strftime('%d.%m.%Y') if self.report.created_at else ''
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = f"Дата: {date_text}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = self.WHITE
        p2.font.name = self.FONT_NAME

        # Описание
        description = self.report.description or 'Нет описания'
        txBox3 = slide.shapes.add_textbox(Inches(1), Inches(5.2), Inches(11), Inches(1.5))
        tf3 = txBox3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = description
        p3.font.size = Pt(12)
        p3.font.color.rgb = self.WHITE
        p3.font.name = self.FONT_NAME

        self._add_footer(slide, 1)